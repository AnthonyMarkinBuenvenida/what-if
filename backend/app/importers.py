"""CSV / XLSX import with validation.

Imports never change a model by themselves. Parsing produces a report (accepted
rows, rejected rows with reasons, warnings) that the user reviews; applying it
creates a new model version. Invalid values are rejected and listed, never
silently replaced.
"""
from __future__ import annotations

import csv
import io
import math
import re
from collections import defaultdict
from datetime import date, datetime

MAX_ROWS = 50_000

COLUMN_ALIASES: dict[str, dict[str, list[str]]] = {
    "products": {
        "name": ["name", "product", "product_name", "item", "sku_name", "description"],
        "id": ["id", "sku", "product_id", "code"],
        "category": ["category", "department", "type"],
        "price": ["price", "selling_price", "sale_price", "retail_price", "unit_price"],
        "unit_cost": ["unit_cost", "cost", "purchase_cost", "cost_price", "wholesale_cost"],
        "starting_inventory": ["starting_inventory", "stock", "inventory", "on_hand", "stock_on_hand", "quantity_on_hand", "qty"],
        "avg_daily_demand": ["avg_daily_demand", "daily_demand", "average_daily_sales", "daily_sales"],
        "shelf_life_days": ["shelf_life_days", "shelf_life", "expiry_days", "days_to_expiry"],
        "reorder_point": ["reorder_point", "reorder_level", "min_stock"],
        "order_quantity": ["order_quantity", "reorder_quantity", "order_qty", "case_qty"],
        "lead_time_days": ["lead_time_days", "lead_time", "supplier_lead_time"],
        "supplier": ["supplier", "vendor"],
    },
    "sales": {
        "date": ["date", "day", "sale_date", "transaction_date"],
        "product": ["product", "name", "product_name", "item", "sku", "product_id"],
        "units": ["units", "quantity", "qty", "units_sold", "quantity_sold"],
        "price": ["price", "unit_price", "selling_price"],
        "stock_on_hand": ["stock_on_hand", "on_hand", "inventory", "stock"],
    },
    "expenses": {
        "name": ["name", "expense", "description", "item"],
        "category": ["category", "type"],
        "amount": ["amount", "cost", "value", "monthly_amount"],
        "frequency": ["frequency", "period", "interval"],
    },
}
REQUIRED = {
    "products": ["name", "price", "unit_cost"],
    "sales": ["date", "product", "units"],
    "expenses": ["name", "amount"],
}
NUMERIC = {
    "products": ["price", "unit_cost", "starting_inventory", "avg_daily_demand", "shelf_life_days", "reorder_point", "order_quantity", "lead_time_days"],
    "sales": ["units", "price", "stock_on_hand"],
    "expenses": ["amount"],
}


class ImportErrorWithReport(Exception):
    def __init__(self, message: str):
        super().__init__(message)
        self.message = message


def _norm_header(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")


def read_table(filename: str, data: bytes) -> list[dict]:
    name = filename.lower()
    if name.endswith(".csv") or name.endswith(".txt"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text), dialect)
        rows = list(reader)
    elif name.endswith(".xlsx") or name.endswith(".xlsm"):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # corrupted or not really an xlsx
            raise ImportErrorWithReport(f"Could not open the spreadsheet: {exc}") from exc
        ws = wb.worksheets[0]
        rows = [["" if c is None else c for c in r] for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif name.endswith(".xls"):
        raise ImportErrorWithReport("Legacy .xls files are not supported. Save the file as .xlsx or .csv and upload again.")
    else:
        raise ImportErrorWithReport("Unsupported file type. Upload a .csv or .xlsx file.")
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        raise ImportErrorWithReport("The file is empty.")
    if len(rows) - 1 > MAX_ROWS:
        raise ImportErrorWithReport(f"The file has {len(rows) - 1:,} rows; the limit is {MAX_ROWS:,}.")
    header = [_norm_header(h) for h in rows[0]]
    return [dict(zip(header, r)) | {"__row__": i + 2} for i, r in enumerate(rows[1:])]


def _map_columns(kind: str, headers: list[str]) -> tuple[dict, list[str]]:
    mapping, unknown = {}, []
    for field, aliases in COLUMN_ALIASES[kind].items():
        for a in aliases:
            if a in headers:
                mapping[field] = a
                break
    used = set(mapping.values())
    unknown = [h for h in headers if h not in used and h != "__row__" and h]
    return mapping, unknown


def _to_number(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if (isinstance(v, float) and math.isnan(v)) else float(v)
    s = str(v).strip().replace("$", "").replace(",", "").replace("€", "").replace("£", "")
    if s == "":
        return None
    if s.endswith("%"):
        s = s[:-1]
    return float(s)


def _to_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%m/%d/%Y", "%d/%m/%Y", "%d.%m.%Y", "%m-%d-%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unrecognised date '{s}'")


def parse_import(kind: str, filename: str, data: bytes, known_products: list[dict] | None = None) -> dict:
    if kind not in COLUMN_ALIASES:
        raise ImportErrorWithReport("Choose products, sales or expenses as the import type.")
    rows = read_table(filename, data)
    headers = [h for h in rows[0].keys() if h != "__row__"]
    mapping, unknown = _map_columns(kind, headers)
    missing = [f for f in REQUIRED[kind] if f not in mapping]
    report = {"kind": kind, "filename": filename, "total_rows": len(rows), "column_mapping": mapping,
              "ignored_columns": unknown, "missing_columns": missing, "errors": [], "warnings": [], "accepted_rows": 0}
    if missing:
        expected = {f: COLUMN_ALIASES[kind][f][:3] for f in missing}
        report["errors"].append({"row": None, "message": "Missing required column(s): " + ", ".join(f"{f} (e.g. {', '.join(expected[f])})" for f in missing)})
        return {"report": report, "parsed": {}}

    clean: list[dict] = []
    for r in rows:
        rec, row_errors = {"__row__": r["__row__"]}, []
        for field, col in mapping.items():
            raw = r.get(col, "")
            if field in NUMERIC[kind]:
                try:
                    rec[field] = _to_number(raw)
                except ValueError:
                    row_errors.append(f"{field} '{raw}' is not a number")
                    continue
                if rec[field] is not None and rec[field] < 0:
                    row_errors.append(f"{field} is negative ({rec[field]:g})")
            elif field == "date":
                try:
                    rec[field] = _to_date(raw)
                except ValueError as exc:
                    row_errors.append(str(exc))
            else:
                rec[field] = str(raw).strip() if raw is not None else ""
        for f in REQUIRED[kind]:
            if f not in rec or rec[f] in (None, ""):
                if not any(e.startswith(f) for e in row_errors):
                    row_errors.append(f"{f} is empty")
        if row_errors:
            report["errors"].append({"row": r["__row__"], "message": "; ".join(row_errors)})
        else:
            clean.append(rec)
    report["accepted_rows"] = len(clean)
    report["rejected_rows"] = len(rows) - len(clean)
    parsed = {"products": _products, "sales": _sales, "expenses": _expenses}[kind](clean, report, known_products or [])
    return {"report": report, "parsed": parsed}


def _products(rows: list[dict], report: dict, _known) -> dict:
    out, seen = [], set()
    for r in rows:
        key = (r.get("id") or r["name"]).strip().lower()
        if key in seen:
            report["warnings"].append({"row": r["__row__"], "message": f"Duplicate product '{r['name']}' – only the first occurrence is used."})
            continue
        seen.add(key)
        if r.get("price") is not None and r.get("unit_cost") is not None and r["price"] < r["unit_cost"]:
            report["warnings"].append({"row": r["__row__"], "message": f"'{r['name']}' sells below cost (${r['price']:.2f} < ${r['unit_cost']:.2f}). Kept as entered."})
        if r.get("shelf_life_days") == 0:
            report["warnings"].append({"row": r["__row__"], "message": f"'{r['name']}' has shelf life 0; treated as non-perishable. Leave blank for non-perishable goods."})
            r["shelf_life_days"] = None
        if r.get("avg_daily_demand") is None:
            report["warnings"].append({"row": r["__row__"], "message": f"'{r['name']}' has no daily demand; import sales history or enter it manually."})
        out.append({k: v for k, v in r.items() if k != "__row__"} | {"source_row": r["__row__"]})
    return {"products": out}


def _sales(rows: list[dict], report: dict, known_products: list[dict]) -> dict:
    by_product: dict[str, dict[date, float]] = defaultdict(lambda: defaultdict(float))
    prices: dict[str, list[float]] = defaultdict(list)
    dup_check: dict[tuple, int] = {}
    for r in rows:
        k = (r["product"].lower(), r["date"])
        if k in dup_check:
            report["warnings"].append({"row": r["__row__"], "message": f"Second row for '{r['product']}' on {r['date']} (first on row {dup_check[k]}); units were added together."})
        else:
            dup_check[k] = r["__row__"]
        by_product[r["product"]][r["date"]] += r["units"]
        if r.get("price"):
            prices[r["product"]].append(r["price"])
        if r.get("stock_on_hand") is not None and r["units"] > r["stock_on_hand"]:
            report["warnings"].append({"row": r["__row__"], "message": f"'{r['product']}' sold {r['units']:g} units on {r['date']} but stock on hand was {r['stock_on_hand']:g}. Check the stock count or sales record."})
    if not by_product:
        return {"products": [], "history_days": 0}
    all_dates = sorted({d for m in by_product.values() for d in m})
    start, end = all_dates[0], all_dates[-1]
    days = (end - start).days + 1
    gaps = days - len(all_dates)
    if gaps > 0:
        report["warnings"].append({"row": None, "message": f"{gaps} day(s) between {start} and {end} have no sales rows for any product. They are counted as zero-sales days; remove them from the range if the store was closed."})
    if days < 28:
        report["warnings"].append({"row": None, "message": f"History covers {days} days. At least 28 days gives a more reliable demand estimate."})
    known = {p["name"].lower(): p for p in known_products} | {str(p.get("id", "")).lower(): p for p in known_products}
    summary = []
    for name, series in by_product.items():
        vals = [series.get(start.fromordinal(start.toordinal() + i), 0.0) for i in range(days)]
        mean = sum(vals) / days
        var = sum((v - mean) ** 2 for v in vals) / max(1, days - 1)
        cv = (math.sqrt(var) / mean) if mean > 0 else 0.0
        match = known.get(name.lower())
        summary.append({
            "product": name, "matched_product_id": match["id"] if match else None,
            "days": days, "total_units": sum(vals), "avg_daily_demand": round(mean, 3), "demand_cv": round(min(cv, 3.0), 3),
            "avg_price": round(sum(prices[name]) / len(prices[name]), 4) if prices[name] else None,
        })
        if not match:
            report["warnings"].append({"row": None, "message": f"Sales product '{name}' does not match any product in the model; it will be added as a new product without price or cost unless you add them."})
    report["warnings"].append({"row": None, "message": "Sales history measures sales, not demand: on days a product was out of stock, true demand was probably higher."})
    return {"products": summary, "history_days": days, "period": [start.isoformat(), end.isoformat()]}


def _expenses(rows: list[dict], report: dict, _known) -> dict:
    out = []
    for r in rows:
        freq = (r.get("frequency") or "monthly").strip().lower()
        freq = {"month": "monthly", "monthly": "monthly", "week": "weekly", "weekly": "weekly", "day": "daily", "daily": "daily"}.get(freq)
        if freq is None:
            report["warnings"].append({"row": r["__row__"], "message": f"Unknown frequency '{r.get('frequency')}' for '{r['name']}'. Row skipped; use daily, weekly or monthly."})
            continue
        cat = (r.get("category") or "").strip().lower()
        if not cat:
            low = r["name"].lower()
            cat = "rent" if "rent" in low or "lease" in low else "utilities" if any(w in low for w in ("electric", "utilit", "water", "gas")) else "other"
        out.append({"name": r["name"], "amount": r["amount"], "frequency": freq, "category": cat, "source_row": r["__row__"]})
    return {"expenses": out}


def apply_import_to_spec(spec: dict, kind: str, parsed: dict) -> tuple[dict, list[str]]:
    """Merge reviewed import data into a copy of the spec. Returns (new spec, change log)."""
    import copy

    from .domain.retail import _slug
    s = copy.deepcopy(spec)
    prov = s.setdefault("provenance", {})
    changes: list[str] = []
    products = s.setdefault("products", [])
    by_key = {p["name"].lower(): p for p in products} | {p["id"].lower(): p for p in products}
    if kind == "products":
        for row in parsed.get("products", []):
            key = (row.get("id") or row["name"]).lower()
            target = by_key.get(key) or by_key.get(row["name"].lower())
            if target is None:
                target = {"id": _slug(row.get("id") or row["name"]), "name": row["name"], "category": row.get("category") or "General",
                          "avg_daily_demand": 0, "demand_cv": 0.3, "starting_inventory": 0, "shelf_life_days": None,
                          "reorder_point": 0, "order_quantity": 0, "lead_time_days": 3, "price_elasticity": -1.0, "price": 0, "unit_cost": 0}
                ids = {p["id"] for p in products}
                while target["id"] in ids:
                    target["id"] += "-x"
                products.append(target)
                by_key[target["name"].lower()] = target
                changes.append(f"Added product {row['name']}")
            for f in ("price", "unit_cost", "starting_inventory", "avg_daily_demand", "shelf_life_days", "reorder_point", "order_quantity", "lead_time_days"):
                if f in row and row[f] is not None:
                    old = target.get(f)
                    target[f] = row[f]
                    prov[f"products.{target['id']}.{f}"] = "confirmed"
                    if old != row[f]:
                        changes.append(f"{target['name']}: {f.replace('_', ' ')} {old} → {row[f]}")
            if row.get("category"):
                target["category"] = row["category"]
    elif kind == "sales":
        for row in parsed.get("products", []):
            target = by_key.get((row.get("matched_product_id") or "").lower()) or by_key.get(row["product"].lower())
            if target is None:
                target = {"id": _slug(row["product"]), "name": row["product"], "category": "General", "price": row.get("avg_price") or 0,
                          "unit_cost": 0, "starting_inventory": 0, "shelf_life_days": None, "reorder_point": 0, "order_quantity": 0,
                          "lead_time_days": 3, "price_elasticity": -1.0}
                products.append(target)
                by_key[target["name"].lower()] = target
                changes.append(f"Added product {row['product']} from sales history (price/cost need review)")
                if row.get("avg_price"):
                    prov[f"products.{target['id']}.price"] = "inferred"
            for f in ("avg_daily_demand", "demand_cv"):
                old = target.get(f)
                target[f] = row[f]
                prov[f"products.{target['id']}.{f}"] = "inferred"
                changes.append(f"{target['name']}: {f.replace('_', ' ')} {old} → {row[f]} (from {row['days']} days of sales)")
        s["sales_history_days"] = parsed.get("history_days")
    elif kind == "expenses":
        expenses = s.setdefault("expenses", [])
        by_name = {e["name"].lower(): e for e in expenses}
        for row in parsed.get("expenses", []):
            target = by_name.get(row["name"].lower())
            if target is None:
                target = {"id": _slug(row["name"]), "name": row["name"]}
                ids = {e["id"] for e in expenses}
                while target["id"] in ids:
                    target["id"] += "-x"
                expenses.append(target)
                by_name[row["name"].lower()] = target
                changes.append(f"Added expense {row['name']}")
            old = target.get("amount")
            target.update({"amount": row["amount"], "frequency": row["frequency"], "category": row["category"]})
            prov[f"expenses.{target['id']}.amount"] = "confirmed"
            if old is not None and old != row["amount"]:
                changes.append(f"{row['name']}: {old} → {row['amount']} ({row['frequency']})")
    return s, changes
